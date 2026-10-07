import csv
from pathlib import Path
from openpyxl import load_workbook

RAW = Path("data/raw")

for file in RAW.iterdir():
    if not file.is_file():
        continue

    print("\n" + "=" * 80)
    print("FILE:", file.name)

    try:
        if file.suffix.lower() == ".csv":
            with open(file, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)

            if not rows:
                print("Empty file")
                continue

            headers = rows[0]
            data = rows[1:]

            print("Shape:", (len(data), len(headers)))
            print("\nColumns:")
            for col in headers:
                print(" -", repr(col))

            missing = []
            for i, col in enumerate(headers):
                count = sum(
                    1 for row in data
                    if i >= len(row) or row[i].strip() == ""
                )
                missing.append((col, count))

            print("\nMissing values:")
            for col, count in missing:
                if count:
                    print(f" - {col}: {count}")

            print("\nFirst 5 rows:")
            for row in data[:5]:
                print(row)

        elif file.suffix.lower() == ".xlsx":
            wb = load_workbook(file, read_only=True, data_only=True)
            ws = wb.active

            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                print("Empty workbook")
                continue

            headers = [str(x) if x is not None else "" for x in rows[0]]
            data = rows[1:]

            print("Sheet:", ws.title)
            print("Shape:", (len(data), len(headers)))

            print("\nColumns:")
            for col in headers:
                print(" -", repr(col))

            print("\nMissing values:")
            for i, col in enumerate(headers):
                count = sum(
                    1 for row in data
                    if i >= len(row) or row[i] is None or str(row[i]).strip() == ""
                )
                if count:
                    print(f" - {col}: {count}")

            print("\nFirst 5 rows:")
            for row in data[:5]:
                print(row)

            wb.close()

    except Exception as e:
        print("ERROR:", type(e).__name__, e)
