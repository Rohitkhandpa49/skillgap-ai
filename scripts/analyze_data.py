import csv
from pathlib import Path
from collections import Counter

RAW = Path("data/raw")

# ---------- CSV READER ----------
def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        rows = list(csv.DictReader(f))
    return rows


# ---------- ANALYTICS JOBS ----------
jobs = read_csv(RAW / "Analytics Jobs.csv")

print("\n" + "="*80)
print("ANALYTICS JOBS — MARKET INTELLIGENCE")
print("="*80)

print("Total jobs:", len(jobs))

# Job titles
titles = Counter(
    r["job_desig"].strip()
    for r in jobs
    if r["job_desig"] and r["job_desig"].strip()
)

print("\nTop 15 job designations:")
for title, count in titles.most_common(15):
    print(f"{title}: {count}")

# Locations
locations = Counter()

for r in jobs:
    value = r["location"].strip() if r["location"] else ""
    if value:
        for location in value.split(","):
            locations[location.strip()] += 1

print("\nTop 15 locations:")
for location, count in locations.most_common(15):
    print(f"{location}: {count}")

# Experience
experience = Counter(
    r["experience"].strip()
    for r in jobs
    if r["experience"] and r["experience"].strip()
)

print("\nTop experience ranges:")
for exp, count in experience.most_common(15):
    print(f"{exp}: {count}")

# Salary bands
salary = Counter(
    r["salary"].strip()
    for r in jobs
    if r["salary"] and r["salary"].strip()
)

print("\nSalary bands:")
for band, count in salary.most_common():
    print(f"{band}: {count}")


# ---------- DATA SCIENCE JOBS ----------
ds = read_csv(RAW / "DataScience Jobs.csv")

print("\n" + "="*80)
print("DATA SCIENCE JOBS — SALARY INTELLIGENCE")
print("="*80)

print("Total records:", len(ds))

companies = Counter(
    r["company_name"].strip()
    for r in ds
    if r["company_name"]
)

print("\nTop companies:")
for company, count in companies.most_common(15):
    print(f"{company}: {count}")

titles = Counter(
    r["job_title"].strip()
    for r in ds
    if r["job_title"]
)

print("\nTop job titles:")
for title, count in titles.most_common(15):
    print(f"{title}: {count}")

# Convert Lakh salary strings
def salary_lakh(value):
    if not value:
        return None
    value = value.strip().upper().replace("L", "")
    try:
        return float(value)
    except:
        return None

avg_salaries = [
    salary_lakh(r["avg_salary"])
    for r in ds
]

avg_salaries = [x for x in avg_salaries if x is not None]

print("\nAverage salary statistics:")
print("Minimum:", min(avg_salaries), "LPA")
print("Maximum:", max(avg_salaries), "LPA")
print("Mean:", round(sum(avg_salaries) / len(avg_salaries), 2), "LPA")


# ---------- JDS SKILL TRAITS ----------
print("\n" + "="*80)
print("JDS SKILL TRAITS")
print("="*80)

try:
    from openpyxl import load_workbook

    wb = load_workbook(
        RAW / "JDS Skill Traits.xlsx",
        read_only=True,
        data_only=True
    )

    ws = wb.active
    rows = list(ws.values)

    headers = list(rows[0])
    data = [dict(zip(headers, row)) for row in rows[1:]]

    # Remove incomplete rows
    data = [
        r for r in data
        if r["salary_hike_high_or_low"] is not None
    ]

    print("Complete records:", len(data))

    target = Counter(
        int(r["salary_hike_high_or_low"])
        for r in data
    )

    print("Salary hike classes:", dict(target))

    features = [
        "big_data_skills",
        "maths-stats_skills",
        "coding_skills",
        "ai_and_ml_skills",
        "dashboard_and_storytelling_skills"
    ]

    for feature in features:
        high = [
            float(r[feature])
            for r in data
            if int(r["salary_hike_high_or_low"]) == 1
        ]

        low = [
            float(r[feature])
            for r in data
            if int(r["salary_hike_high_or_low"]) == 0
        ]

        print(
            f"{feature}: "
            f"High={sum(high)/len(high):.2f}, "
            f"Low={sum(low)/len(low):.2f}"
        )

    wb.close()

except Exception as e:
    print("JDS analysis error:", e)


# ---------- SDS PERSONALITY ----------
print("\n" + "="*80)
print("SDS PERSONALITY TRAITS")
print("="*80)

try:
    wb = load_workbook(
        RAW / "SDS Personality Traits.xlsx",
        read_only=True,
        data_only=True
    )

    ws = wb.active
    rows = list(ws.values)

    headers = list(rows[0])
    data = [dict(zip(headers, row)) for row in rows[1:]]

    target_name = "success_ classification_ high_low"

    data = [
        r for r in data
        if r[target_name] is not None
    ]

    print("Complete records:", len(data))

    target = Counter(
        int(r[target_name])
        for r in data
    )

    print("Success classes:", dict(target))

    features = [
        "neuroticism",
        " extraversion",
        "openness_to_experience",
        "agreeableness",
        "conscientiousness"
    ]

    for feature in features:
        high = [
            float(r[feature])
            for r in data
            if int(r[target_name]) == 1
        ]

        low = [
            float(r[feature])
            for r in data
            if int(r[target_name]) == 0
        ]

        print(
            f"{feature.strip()}: "
            f"High={sum(high)/len(high):.2f}, "
            f"Low={sum(low)/len(low):.2f}"
        )

    wb.close()

except Exception as e:
    print("SDS analysis error:", e)
