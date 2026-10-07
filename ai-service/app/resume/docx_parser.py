"""DOCX Document Text Extraction using python-docx.

Extracts text from Microsoft Word documents safely:
- Extracts content from standard paragraphs
- Extracts content from tables and cells
- Ignores embedded macros or scripts (treats XML/Word solely as static document data)
- Tracks paragraph and table counts
- Handles corrupt or invalid DOCX files safely with CorruptDocxError
"""

from pathlib import Path
from typing import Dict, List, Tuple
from docx import Document

from app.resume.exceptions import CorruptDocxError


def extract_text_from_docx(docx_path: Path) -> Tuple[str, Dict[str, int]]:
    """Extract raw text, paragraphs, and tables from a DOCX file.

    Args:
        docx_path: Path to target DOCX document on disk.

    Returns:
        Tuple of (raw_concatenated_text, metadata_dict)
        where metadata_dict contains:
            paragraph_count, table_count, character_count

    Raises:
        CorruptDocxError: If the DOCX cannot be opened or parsed.
    """
    text_blocks: List[str] = []

    try:
        doc = Document(docx_path)
        paragraph_count = len(doc.paragraphs)
        table_count = len(doc.tables)

        # 1. Extract paragraphs
        for p in doc.paragraphs:
            stripped = p.text.strip()
            if stripped:
                text_blocks.append(stripped)

        # 2. Extract tables
        for table in doc.tables:
            for row in table.rows:
                # Deduplicate cells in case of merged cells within the same row
                seen_cells: List[str] = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in seen_cells:
                        seen_cells.append(cell_text)
                if seen_cells:
                    text_blocks.append(" | ".join(seen_cells))

    except Exception as ex:
        raise CorruptDocxError(f"Unable to read DOCX file '{docx_path.name}': {str(ex)}") from ex

    concatenated_raw_text = "\n\n".join(text_blocks)

    metadata = {
        "paragraph_count": paragraph_count,
        "table_count": table_count,
        "character_count": len(concatenated_raw_text),
    }

    return concatenated_raw_text, metadata
